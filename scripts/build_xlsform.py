"""Generate a Survey123 XLSForm from docs/schema.yaml.

Usage: python build_xlsform.py docs/schema.yaml form/whatcom_salmon.xlsx [--version 2026100101]
"""
import argparse
import datetime as dt
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font

SURVEY_COLS = ["type", "name", "label", "hint", "required", "required_message",
               "constraint", "constraint_message", "relevant", "default",
               "calculation", "choice_filter", "appearance", "bind::esri:fieldType",
               "bind::esri:fieldLength", "body::accuracyThreshold"]

TYPE_MAP = {"datetime": "dateTime", "text": "text", "geopoint": "geopoint",
            "integer": "integer", "decimal": "decimal", "image": "image"}


def condition_expr(cond: dict) -> str:
    parts = []
    for field, values in cond.items():
        parts += [f"${{{field}}} = '{v}'" for v in values]
    return " or ".join(parts)


def survey_row(f: dict, schema: dict) -> dict:
    r = {c: "" for c in SURVEY_COLS}
    r["name"], r["label"] = f["name"], f["label"]
    t = f["type"]
    r["type"] = f"select_one {f['domain']}" if t == "select_one" else TYPE_MAP[t]

    if f.get("system"):
        r["type"] = "hidden"
        r["bind::esri:fieldType"] = ("esriFieldTypeDouble" if t == "decimal"
                                     else "esriFieldTypeString")
    if f.get("calc") == "gps_accuracy":
        r["calculation"] = 'pulldata("@geopoint", ${location}, "horizontalAccuracy")'

    if f.get("required"):
        r["required"] = "yes"
        r["required_message"] = f"{f['label']} is required"
    if "required_if" in f:
        r["relevant"] = condition_expr(f["required_if"])
        r["required"] = "yes"  # only enforced while relevant
        r["required_message"] = f"{f['label']} is required for this observation type"

    cons, msgs = [], []
    if "min" in f and "max" in f:
        cons.append(f". >= {f['min']} and . <= {f['max']}")
        msgs.append(f"Must be {f['min']}-{f['max']}")
    if "max_length" in f:
        r["bind::esri:fieldLength"] = f["max_length"]
        if not f.get("system"):
            cons.append(f"string-length(.) <= {f['max_length']}")
            msgs.append(f"Max {f['max_length']} characters")
    if f.get("not_future"):
        cons.append(". <= now()")
        msgs.append("Date can't be in the future")
    r["constraint"] = " and ".join(f"({c})" for c in cons) if len(cons) > 1 else "".join(cons)
    r["constraint_message"] = "; ".join(msgs)

    if f.get("filter_by"):
        r["choice_filter"] = f"{f['filter_by']}=${{{f['filter_by']}}}"   # cascading select
    if t == "geopoint":
        r["body::accuracyThreshold"] = schema.get("gps_accuracy_block_m", schema["gps_accuracy_warn_m"])
    if f["name"] == "observer":
        r["calculation"] = "${username}"   # VERIFY: editable default from signed-in user
    if f["name"] == "notes":
        r["appearance"] = "multiline"
    return r


def build(schema_path: str, out_path: str, version: str) -> None:
    schema = yaml.safe_load(open(schema_path))
    wb = Workbook()

    ws = wb.active
    ws.title = "survey"
    ws.append(SURVEY_COLS)
    ws.append(["username", "username"] + [""] * (len(SURVEY_COLS) - 2))
    for f in schema["fields"]:
        r = survey_row(f, schema)
        ws.append([r[c] for c in SURVEY_COLS])

    ch = wb.create_sheet("choices")
    filters = schema.get("domain_filters", {})
    fcols = sorted({v["by"] for v in filters.values()})
    ch.append(["list_name", "name", "label"] + fcols)
    for dname, codes in schema["domains"].items():
        fmap = filters.get(dname, {})
        for code, label in codes.items():
            ch.append([dname, code, label] + [fmap["map"][code] if fmap.get("by") == c else ""
                                              for c in fcols])

    st = wb.create_sheet("settings")
    st.append(["form_title", "form_id", "version", "instance_name"])
    st.append(["Whatcom Creek Salmon Survey", schema["dataset"].replace("-", "_"),
               version, "concat(${station}, ' - ', ${obs_type})"])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    wb.save(out_path)
    print(f"Wrote {out_path} ({len(schema['fields'])} fields, version {version})")


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("schema")
    p.add_argument("out")
    p.add_argument("--version", default=dt.datetime.now().strftime("%Y%m%d%H"))
    a = p.parse_args()
    build(a.schema, a.out, a.version)
